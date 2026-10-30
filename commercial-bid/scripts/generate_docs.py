# -*- coding: utf-8 -*-
"""根据配置 JSON 生成政府采购/招投标商务标 Word 草稿。

用法:
    python generate_docs.py --config config.json --out 输出目录

配置模板见同目录 config.example.json。
"""
import argparse
import json
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def load_config(path):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def set_font(run, size=12, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold


def new_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2)
    sec.left_margin = Cm(2)
    sec.right_margin = Cm(2)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    st.font.size = Pt(12)
    return doc


def title(doc, text, size=16):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(text), size=size, bold=True)
    return p


def para(doc, text, bold=False):
    p = doc.add_paragraph()
    set_font(p.add_run(text), bold=bold)
    return p


def save(doc, out_dir, name):
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, name)
    doc.save(path)
    print("saved", path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--config", required=True)
    ap.add_argument("--out", default="商务标草稿")
    args = ap.parse_args()

    c = load_config(args.config)
    p = c["project"]
    s = c["supplier"]
    sme = s.get("sme", {})
    agent = s.get("agent", {})

    date = p.get("bid_opening_date", "【待填】")
    date_cn = p.get("bid_opening_date_cn", "【待填】")
    project = p["name"]
    project_no = p.get("no", "【待填】")
    buyer = p.get("buyer", "【待填】")
    agency = p.get("agency", "【待填】")
    company = s["company"]
    legal = s.get("legal_rep", "【待填】")
    legal_title = s.get("legal_rep_title", "经理")
    address = s.get("address", "【待填】")
    phone = s.get("phone", "【待填】")
    zipcode = s.get("zip", "【待填】")
    credit = s.get("credit_code", "【待填】")
    bank = s.get("bank", "【待填】")
    account = s.get("account", "【待填】")
    nature = s.get("nature", "【待填：以营业执照为准】")
    founded = s.get("founded", "【待填】")
    capital = s.get("registered_capital", "【待填】")
    scope = s.get("scope", "【待填：以营业执照载明为准】")
    guarantee = p.get("bid_guarantee", "0")
    guarantee_note = p.get("bid_guarantee_note", "本项目不收取投标保证金")
    service_days = p.get("service_days", "365")
    service_period = p.get("service_period", "一年")
    validity = p.get("bid_validity_days", "50")
    price_cn = p.get("price_total_cn", "【待填：人民币____元整】")
    price_num = p.get("price_total", "【待填】")
    industry = sme.get("industry", "物业管理")
    employees = sme.get("employees", "【待填】")
    revenue = sme.get("revenue", "【待填】")
    assets = sme.get("assets", "【待填】")
    size = sme.get("size", "【待填】")
    agent_name = agent.get("name", "【待填】")
    agent_title = agent.get("title", "【待填】")
    fixed_phone = s.get("fixed_phone", "【待填】")

    # 格式一、投标函
    d = new_doc()
    title(d, "投标函")
    para(d, f"{agency}：")
    para(d, f"根据你方政府采购 {project} 的项目编号为：{project_no} 的招标文件，我方正式授权的下述签字人 {legal}（职务：{legal_title}）代表我方 {company}，按照你方招标文件的规定，提交全部文件正本1份、副本【待定】份，并保证所提供的全部文件是真实的、有效的和准确的。")
    para(d, "据此函，签字人兹宣布同意如下：")
    para(d, f"1.按招标文件规定提供服务及服务的投标总价为（大写）{price_cn}元人民币。其中分包报价为：本项目不分包。")
    para(d, "2.如果我方中标，我们保证根据招标文件规定履行合同责任和义务。具体服务期限承诺如下：")
    para(d, f"合同订立后 {service_days} 天；自【待定】年【待定】月【待定】日起至【待定】年【待定】月【待定】日止，期限为{service_period}。")
    para(d, f"3.我方人民币 {guarantee} 元的投标保证金与本投标文件同时提交（{guarantee_note}）。")
    para(d, "4.如果我方中标，我方保证按照招标文件规定提交履约保证金，承担履约责任。")
    para(d, "5.我们已详细阅读了全部招标文件，包括招标文件的修改、补充文件、参考资料及有关的附件，我们接受招标文件的全部条款和条件，我们知道必须放弃提出含糊不清或误解的问题的权利。")
    para(d, "6.我们对招标文件关于时限、程序方面的规定没有异议，保证按照招标文件规定的时限和程序参加投标活动。")
    para(d, "7.我们同意在投标人须知规定的开标时间起遵循本投标书，并在投标人须知规定的投标有效期满之前均具有约束力，并有可能中标。")
    para(d, "8.我们如果在规定的投标有效期内撤回投标，则你方可不予退还我们的投标保证金。")
    para(d, "9.我们保证向你方提供你方可能要求的与本投标有关的任何证据或资料。")
    para(d, "10.我们完全理解你方不一定要接受最低报价的投标或收到的任何投标。")
    para(d, f"11.本投标自开标之时起{validity}天内有效。")
    para(d, "12.我方保证严格遵守《中华人民共和国政府采购法》的有关规定，若有下列情形之一的，我方将被处不予退还投标保证金，列入不良行为记录名单，在一至三年内禁止参加政府采购活动，并予以公告：")
    for item in [
        "（1）提供虚假材料谋取中标、成交的；",
        "（2）采取不正当手段诋毁、排挤其他投标人的；",
        "（3）与其他投标人、采购中心或者采购单位人员或者其他有关人员恶意串通的；",
        "（4）向评标委员会成员、采购中心或者采购单位人员或者其他有关人员行贿或者提供其他不正当利益的；",
        "（5）被评定中标后无正当理由不与采购人或者采购代理机构订立合同，或者中标后不按招标文件和中标供应商的投标文件订立合同，或者与采购人另行订立背离合同实质性内容的协议的；",
        "（6）将中标项目转让给他人或者将中标项目分包给他人的；",
        "（7）签订合同后拒绝履行合同义务的；",
        "（8）拒绝有关部门监督检查或提供虚假情况的。",
    ]:
        para(d, item)
    para(d, "")
    para(d, f"投标人名称（加盖公章）：{company}")
    para(d, f"地址：{address}")
    para(d, f"电话、传真或电传：{phone}")
    para(d, f"邮政编码：{zipcode}")
    para(d, f"日期：{date}")
    save(d, args.out, "01投标函.docx")

    # 格式二、开标一览表
    d = new_doc()
    title(d, "开标一览表")
    para(d, f"项目名称：{project}")
    para(d, f"项目编号：{project_no}")
    para(d, f"投标人名称（加盖投标人公章）：{company}")
    para(d, date)
    table = d.add_table(rows=2, cols=5)
    table.style = "Table Grid"
    for i, h in enumerate(["投标包号", "投标总价（元）", "服务期限", "投标保证金（元）", "备注"]):
        table.rows[0].cells[i].text = h
    for i, v in enumerate(["/", f"【待填：大写____元；小写{price_num}元】", service_period, guarantee, guarantee_note]):
        table.rows[1].cells[i].text = v
    para(d, "")
    para(d, "投标要求：")
    para(d, "1.“开标一览表”用于开标时唱标使用。")
    para(d, "2、如果给予价格折扣，必须在“开标一览表”中填报，否则，不作为评标依据。")
    para(d, "3、“开标一览表”的内容应与“投标报价明细表”以及投标文件的其他相关内容一致。")
    para(d, "4、“开标一览表”中各个栏目都必须完整、准确填写。开标时，“开标一览表”的所有内容都不允许补充或者修改。")
    save(d, args.out, "02开标一览表.docx")

    # 格式三、投标人基本情况
    d = new_doc()
    title(d, "投标人基本情况")
    para(d, f"企业性质：{nature}")
    para(d, f"企业名称：{company}")
    para(d, f"统一社会信用代码：{credit}")
    para(d, f"法定代表人：{legal}")
    para(d, f"注册地址：{address}")
    para(d, f"成立日期：{founded}")
    para(d, f"注册资本：{capital}")
    para(d, f"经营范围：{scope}")
    para(d, "")
    para(d, "发展历程与经营规模：公司长期从事物业管理、环境卫生管理及建筑物清洁服务，具备机关、学校、医院等场所保洁及物业服务的现场管理经验，服务团队稳定。")
    para(d, "服务理念：坚持以客户满意为中心，规范作业、快速响应、持续改进，严格落实院感防控与安全生产要求。")
    para(d, f"主营服务：{scope}")
    para(d, "技术力量：公司建立人员培训、岗位考核、院感防控、安全作业、设备管理、质量巡检等制度，配备高压清洗机、驾驶式洗地机、尘推车、单擦机、吸水机等保洁设备，能够保障本项目保洁服务稳定运行。")
    para(d, "")
    para(d, "（此处可附主要经营场所、主要服务、服务流程等图片说明）")
    save(d, args.out, "03投标人基本情况.docx")

    # 格式四、投标人的资格声明
    d = new_doc()
    title(d, "投标人的资格声明")
    para(d, f"{agency}：")
    para(d, f"按照《中华人民共和国政府采购法》第22条和你中心发布的 {project} 采购项目招标文件》（编号：{project_no}）的规定，我公司郑重声明如下：")
    para(d, f"1、我公司是按照中华人民共和国法律在工商管理机关登记注册的企业法人，注册地点为 {address}，公司全称为 {company}，法定代表人为 {legal}，具有独立承担民事责任的能力。")
    para(d, "2、我公司具有良好的商业信誉和健全的财务会计制度。")
    para(d, "3、我公司具有履行本项目采购合同所必需的设备和专业技术能力。")
    para(d, "4、我公司具有依法缴纳税收和社会保障资金的良好记录。")
    para(d, "5、我公司在参与本次政府采购活动之前三年内，在经营活动中无重大违法记录。")
    para(d, "6、我公司在参与本次政府采购活动时未受到任何地方政府采购部门做出的暂停参加政府采购活动的处罚。")
    para(d, "我公司保证上述声明的事项都是真实的，如有虚假，我公司承担相应的法律责任，并承担因此给你中心以及本项目采购人所造成的损失。")
    para(d, "我公司已经按照你中心招标文件的要求提交了所要求提交的能够证明上述声明事项真实性的全部文件材料，并保证随时按照你中心的要求提供能够证明上述声明事项真实性的任何有效文件。")
    para(d, "")
    para(d, f"投标人全称并加盖单位公章：{company}")
    para(d, f"时间：{date_cn}")
    save(d, args.out, "04投标人的资格声明.docx")

    # 无重大违法记录声明
    d = new_doc()
    title(d, "无重大违法记录声明")
    para(d, f"致：{agency}、{buyer}")
    para(d, f"我公司 {company} 郑重声明：")
    para(d, "我公司在参加本次政府采购活动前三年内（声明出具之日前36个月），在经营活动中没有重大违法记录，即未因违法经营受到刑事处罚，未受到责令停产停业、吊销许可证或者执照、较大数额罚款等行政处罚，也不存在因违法经营被禁止在一定期限内参加政府采购活动而期限尚未届满的情形。")
    para(d, "本声明内容真实、合法、有效，如有虚假，我公司愿承担相应法律责任，并赔偿由此给采购人及采购代理机构造成的损失。")
    para(d, "")
    para(d, f"声明人（加盖公章）：{company}")
    para(d, f"日期：{date}")
    save(d, args.out, "05无重大违法记录声明.docx")

    # 格式六、中小企业声明函
    d = new_doc()
    title(d, "中小企业声明函（工程、服务）")
    para(d, f"本公司（联合体）郑重声明，根据《政府采购促进中小企业发展管理办法》（财库﹝2020﹞46 号）的规定，本公司（联合体）参加（{buyer}）的（{project}）采购活动，服务全部由符合政策要求的中小企业承接。相关企业（含联合体中的中小企业、签订分包意向协议的中小企业）的具体情况如下：")
    para(d, f"1. （标的名称）保洁服务，属于（采购文件中明确的所属行业）{industry}行业；承建（承接）企业为（企业名称）{company}，从业人员 {employees} 人，营业收入为 {revenue} 万元，资产总额为 {assets} 万元，属于（{size}）；")
    para(d, "以上企业，不属于大企业的分支机构，不存在控股股东为大企业的情形，也不存在与大企业的负责人为同一人的情形。")
    para(d, "本企业对上述声明内容的真实性负责。如有虚假，将依法承担相应责任。")
    para(d, "")
    para(d, f"企业名称（盖章）：{company}")
    para(d, f"日期：{date}")
    para(d, "")
    para(d, "注：从业人员、营业收入、资产总额填报上一年度数据，无上一年度数据的新成立企业可不填报。")
    save(d, args.out, "06中小企业声明函.docx")

    # 格式七、售后服务承诺书
    d = new_doc()
    title(d, "售后服务承诺书")
    para(d, f"{agency}及采购单位（{buyer}）：")
    para(d, f"我公司 {company} 自愿参加 {buyer} 组织的 {project}（招标编号：{project_no}）的投标。我公司郑重承诺，如果我公司的投标被评定为中标，我公司对于中标服务，除完全响应招标文件对伴随服务和售后服务的所有要求外，还将按照以下条款提供优质和完善的售后服务：")
    para(d, "1、我公司中标后将为采购人提供下列售后服务项目：")
    para(d, "（1）按合同约定持续提供保洁、消毒及院感防控等全部服务；")
    para(d, "（2）配合采购人等级评审、大型检查及突发公共事件，无条件增派人员、强化清洁消毒；")
    para(d, "（3）定期开展人员培训、设备维护、耗材补充，保障服务连续性。")
    para(d, "2、我公司的售后服务响应及到达现场的时间：")
    para(d, "（1）一般诉求2小时内响应、24小时内整改完毕；")
    para(d, "（2）紧急污染、地面泼洒、设施损坏等突发情况30分钟内到场处置；")
    para(d, "（3）接到通知后应急增派人员1小时内到位；设备故障2小时内响应、24小时内修复或调配备用设备。")
    para(d, "3、我公司对本项目的技术培训安排：")
    para(d, "（1）进场前完成岗前培训，内容包括服务意识、岗位技能、院感知识、操作规范、安全防护、消防及应急预案，考核合格后上岗；")
    para(d, "（2）每季度组织集中培训并考核，不合格人员暂停上岗、补考合格后复工。")
    para(d, "4、我公司用于本项目的维修技术人员及设备、备品、备件供应的保证措施及收费标准：")
    para(d, "我公司按招标文件配备保洁设备、工具与耗材，建立设备维护台账，故障设备及时维修或调配备用，确保作业不中断；上述费用已包含在投标总价内，不另行收费。")
    para(d, "5、我公司在本项目所在地服务网点明细表及相关情况：")
    para(d, "我公司在本项目所在地设立的分支服务机构一览表")
    t2 = d.add_table(rows=2, cols=6)
    t2.style = "Table Grid"
    for i, h in enumerate(["序号", "分支服务机构名称", "所在市县及街区门牌号", "联系人", "移动电话", "固定电话"]):
        t2.rows[0].cells[i].text = h
    for i, v in enumerate(["1", f"{company}（驻地服务网点）", address, legal, phone, fixed_phone]):
        t2.rows[1].cells[i].text = v
    para(d, "")
    para(d, f"投标人全称（加盖投标人公章）：{company}")
    para(d, date)
    save(d, args.out, "07售后服务承诺书.docx")

    # 格式五、法定代表人授权书（可选）
    d = new_doc()
    title(d, "法定代表人授权书（法定代表人投标的可不提供）")
    para(d, f"本授权书声明：注册于（注册地点）{address} 的 {company} 公司的（投标人法定代表人姓名、职务）{legal}（{legal_title}）代表本公司授权___（被授权人单位名称）{company}__的___（被授权人姓名、职务）{agent_name}（{agent_title}）为本公司的合法代理人，就白城市政府采购项目（招标编号：{project_no}）的投标以及合同的谈判、签约、执行、完成和保修等全权负责，以本公司名义处理一切与之有关的事务。代理人在投标、开标、评标、合同谈判和履行过程中所签署的一切文件和处理的与之有关的一切事务，我均予以承认。代理人无转委托权。")
    para(d, f"本授权书于 {date} 签字生效，特此声明。")
    para(d, "")
    para(d, f"投标人名称（加盖公章）：{company}")
    para(d, f"地址：{address}")
    para(d, f"授权人（法定代表人）姓名：{legal}")
    para(d, f"被授权人（代理人）姓名：{agent_name}")
    para(d, "")
    para(d, "请投标人注意：法定代表人授权书必须按规定有效签署和加盖公章。")
    save(d, args.out, "08法定代表人授权书（可选）.docx")

    print("done")


if __name__ == "__main__":
    main()
